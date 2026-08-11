import os
import sys
from reportlab.lib.pagesizes import landscape
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# CONSTANTS & METADATA FOR WEEK 9
# -----------------------------------------------------------------------------
INTERN_ID = "TBI-26100454"
INTERN_NAME = "Sanchi Sisodia"
UNIVERSITY = "Graphic Era University (GEU)"
PROJECT_TITLE = "CareerPilot AI"
WEEK_TITLE = "Week 9 — App Deployment & Go-Live"

URL_FRONTEND = "https://careerpilot-ai.vercel.app"
URL_BACKEND = "https://careerpilot-ai-api.onrender.com"
URL_GITHUB = "https://github.com/sanchi1905/careerpilot-ai"

IMG_DASHBOARD = "dashboard_overview.png"
IMG_MODAL = "add_application_modal.png"
IMG_SCHEMA = "schema_diagram.png"
IMG_NETWORK = "network_tab.png"

# Colors
PRIMARY_HEX = "#4F46E5"     # Indigo 600
PRIMARY_DARK = "#3730A3"    # Indigo 800
ACCENT_HEX = "#06B6D4"      # Cyan 500
TEXT_DARK = "#1E293B"       # Slate 800
TEXT_MUTED = "#64748B"      # Slate 500
BG_LIGHT = "#F8FAFC"        # Slate 50
BORDER_HEX = "#E2E8F0"      # Slate 200

COLOR_PRIMARY = colors.HexColor(PRIMARY_HEX)
COLOR_PRIMARY_DARK = colors.HexColor(PRIMARY_DARK)
COLOR_ACCENT = colors.HexColor(ACCENT_HEX)
COLOR_TEXT_DARK = colors.HexColor(TEXT_DARK)
COLOR_TEXT_MUTED = colors.HexColor(TEXT_MUTED)
COLOR_BG_LIGHT = colors.HexColor(BG_LIGHT)
COLOR_BORDER = colors.HexColor(BORDER_HEX)

SLIDE_WIDTH = 960
SLIDE_HEIGHT = 540

class W9NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            self.setFillColor(COLOR_PRIMARY)
            self.rect(0, 0, 24, SLIDE_HEIGHT, fill=1, stroke=0)
            self.setFillColor(COLOR_ACCENT)
            self.rect(24, 0, 8, SLIDE_HEIGHT, fill=1, stroke=0)
            return

        # Top Header Bar
        self.setFillColor(COLOR_PRIMARY)
        self.rect(0, SLIDE_HEIGHT - 12, SLIDE_WIDTH, 12, fill=1, stroke=0)
        
        self.setFont("Helvetica-Bold", 9)
        self.setFillColor(COLOR_TEXT_MUTED)
        self.drawString(40, SLIDE_HEIGHT - 28, f"{PROJECT_TITLE} — {WEEK_TITLE}")
        
        self.setFont("Helvetica", 9)
        self.drawRightString(SLIDE_WIDTH - 40, SLIDE_HEIGHT - 28, f"Intern ID: {INTERN_ID}")
        
        self.setStrokeColor(COLOR_BORDER)
        self.setLineWidth(0.75)
        self.line(40, SLIDE_HEIGHT - 34, SLIDE_WIDTH - 40, SLIDE_HEIGHT - 34)

        # Footer
        self.line(40, 36, SLIDE_WIDTH - 40, 36)
        self.setFont("Helvetica", 8)
        self.setFillColor(COLOR_TEXT_MUTED)
        self.drawString(40, 20, f"{INTERN_NAME} | {UNIVERSITY} | Deployment Proof & Presentation")
        self.drawRightString(SLIDE_WIDTH - 40, 20, f"Slide {self._pageNumber} of {page_count}")


def build_w9_pdf(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=(SLIDE_WIDTH, SLIDE_HEIGHT),
        leftMargin=40,
        rightMargin=40,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'W9Title',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=COLOR_PRIMARY_DARK,
        spaceAfter=12
    )

    badge_style = ParagraphStyle(
        'W9Badge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=COLOR_PRIMARY,
        spaceAfter=4,
        textTransform='uppercase'
    )

    body_style = ParagraphStyle(
        'W9Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=16,
        textColor=COLOR_TEXT_DARK,
        spaceAfter=8
    )

    bullet_style = ParagraphStyle(
        'W9Bullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=15,
        textColor=COLOR_TEXT_DARK,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=6
    )

    bold_b = lambda title, text: f"<b><font color='{PRIMARY_DARK}'>• {title}:</font></b> {text}"

    story = []

    # -------------------------------------------------------------------------
    # SLIDE 1: Title
    # -------------------------------------------------------------------------
    story.append(Spacer(1, 40))
    story.append(Paragraph("WEEK 9 SUBMISSION DELIVERABLE", badge_style))
    story.append(Paragraph(f"<b><font size=28 color='{PRIMARY_HEX}'>{PROJECT_TITLE}</font></b>", title_style))
    story.append(Paragraph(f"<b><font size=14 color='{TEXT_MUTED}'>{WEEK_TITLE}</font></b>", body_style))
    story.append(Spacer(1, 20))

    meta_data = [
        [Paragraph("<b>Intern Name:</b>", body_style), Paragraph(INTERN_NAME, body_style),
         Paragraph("<b>Intern ID:</b>", body_style), Paragraph(f"<font color='{PRIMARY_HEX}'><b>{INTERN_ID}</b></font>", body_style)],
        [Paragraph("<b>University:</b>", body_style), Paragraph(UNIVERSITY, body_style),
         Paragraph("<b>Submission:</b>", body_style), Paragraph("Week 9 Deployment Proof", body_style)],
        [Paragraph("<b>Live Frontend:</b>", body_style), Paragraph(f"<a href='{URL_FRONTEND}' color='{PRIMARY_HEX}'><u>{URL_FRONTEND}</u></a>", body_style),
         Paragraph("<b>Live Backend:</b>", body_style), Paragraph(f"<a href='{URL_BACKEND}' color='{PRIMARY_HEX}'><u>{URL_BACKEND}</u></a>", body_style)]
    ]
    t_meta = Table(meta_data, colWidths=[110, 310, 100, 310])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), COLOR_BG_LIGHT),
        ('BOX', (0,0), (-1,-1), 1, COLOR_BORDER),
        ('INNERGRID', (0,0), (-1,-1), 0.5, COLOR_BORDER),
        ('PADDING', (0,0), (-1,-1), 10),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t_meta)
    story.append(PageBreak())

    # -------------------------------------------------------------------------
    # SLIDE 2: Deployment Overview
    # -------------------------------------------------------------------------
    story.append(Paragraph("DEPLOYMENT OVERVIEW", badge_style))
    story.append(Paragraph("Production Cloud Hosting Architecture", title_style))

    overview_data = [
        [Paragraph("<b>Layer</b>", body_style), Paragraph("<b>Hosting Platform</b>", body_style), Paragraph("<b>Production Configuration</b>", body_style)],
        [
            Paragraph("<b>Frontend UI</b>", body_style),
            Paragraph("<b>Vercel</b>", body_style),
            Paragraph("• Deployed from GitHub repository root (<code>frontend/</code>).<br/>• Configured production environment variables (<code>VITE_API_URL</code>).<br/>• Global Edge CDN hosting with zero downtime.", bullet_style)
        ],
        [
            Paragraph("<b>Backend API</b>", body_style),
            Paragraph("<b>Render</b>", body_style),
            Paragraph("• Node.js web service running <code>npm start</code> on production port.<br/>• Environment variables configured: <code>MONGO_URI</code>, <code>JWT_SECRET</code>, <code>FRONTEND_ORIGIN</code>.<br/>• Native CORS protection linked strictly to Vercel origin.", bullet_style)
        ],
        [
            Paragraph("<b>Database</b>", body_style),
            Paragraph("<b>MongoDB Atlas</b>", body_style),
            Paragraph("• Cloud-hosted M0 Free Cluster (Mongoose ODM).<br/>• IP Access Whitelisted (<code>0.0.0.0/0</code>) for Render cloud instances.<br/>• Persistent data storage surviving server restarts.", bullet_style)
        ]
    ]
    t_ov = Table(overview_data, colWidths=[120, 180, 560])
    t_ov.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), COLOR_PRIMARY),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('BOX', (0,0), (-1,-1), 1, COLOR_BORDER),
        ('INNERGRID', (0,0), (-1,-1), 0.5, COLOR_BORDER),
        ('PADDING', (0,0), (-1,-1), 10),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    for i in range(3):
        overview_data[0][i].style.textColor = colors.white

    story.append(t_ov)
    story.append(PageBreak())

    # -------------------------------------------------------------------------
    # SLIDE 3: Live Application Demonstration
    # -------------------------------------------------------------------------
    story.append(Paragraph("PRODUCTION VERIFICATION", badge_style))
    story.append(Paragraph("Live Application UI & Feature Proof", title_style))

    story.append(Paragraph(f"<b>Live Public URL:</b> <a href='{URL_FRONTEND}' color='{PRIMARY_HEX}'><u>{URL_FRONTEND}</u></a> — All core features active on production.", body_style))
    story.append(Spacer(1, 6))

    if os.path.exists(IMG_DASHBOARD) and os.path.exists(IMG_MODAL):
        img1 = Image(IMG_DASHBOARD, width=425, height=240)
        img2 = Image(IMG_MODAL, width=425, height=240)
        t_imgs = Table([[
            [img1, Paragraph("<b>Figure 1:</b> Live Application Dashboard & Analytics", bullet_style)],
            [img2, Paragraph("<b>Figure 2:</b> Add Application Modal & Data Creation", bullet_style)]
        ]], colWidths=[435, 435])
        t_imgs.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))
        story.append(t_imgs)

    story.append(PageBreak())

    # -------------------------------------------------------------------------
    # SLIDE 4: End-to-End API & AI Verification
    # -------------------------------------------------------------------------
    story.append(Paragraph("API & AI FUNCTIONALITY", badge_style))
    story.append(Paragraph("Live REST API & Google Gemini AI Verification", title_style))

    c1 = [
        Paragraph("<b>Backend API Endpoints</b>", ParagraphStyle('H1', parent=body_style, textColor=COLOR_PRIMARY_DARK)),
        Paragraph(bold_b("GET /api/applications/stats", "Aggregates total applications, interviews, offers, and average resume match score."), bullet_style),
        Paragraph(bold_b("POST /api/applications", "Creates new job application record in MongoDB Atlas with validation."), bullet_style),
        Paragraph(bold_b("POST /api/ai/prep", "Invokes Google Gemini 1.5 Flash to generate 3 role questions + 1 tip."), bullet_style),
    ]

    c2 = [
        Paragraph("<b>Deployment Security & Integrity</b>", ParagraphStyle('H2', parent=body_style, textColor=COLOR_PRIMARY_DARK)),
        Paragraph(bold_b("JWT Authentication", "Stateless JWT headers required for all application & AI endpoints."), bullet_style),
        Paragraph(bold_b("CORS Policy", "Backend dynamically enforces <code>FRONTEND_ORIGIN</code> matching Vercel domain."), bullet_style),
        Paragraph(bold_b("Rate Limiting", "Express rate limiter enabled on authentication endpoints to prevent abuse."), bullet_style),
    ]

    t_api = Table([[c1, c2]], colWidths=[430, 430])
    t_api.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), COLOR_BG_LIGHT),
        ('BACKGROUND', (1,0), (1,0), COLOR_BG_LIGHT),
        ('BOX', (0,0), (0,0), 1, COLOR_BORDER),
        ('BOX', (1,0), (1,0), 1, COLOR_BORDER),
        ('PADDING', (0,0), (-1,-1), 12),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(t_api)
    story.append(Spacer(1, 10))

    if os.path.exists(IMG_NETWORK):
        img_net = Image(IMG_NETWORK, width=550, height=155)
        story.append(img_net)

    story.append(PageBreak())

    # -------------------------------------------------------------------------
    # SLIDE 5: Debugging Deployment & Known Limitations
    # -------------------------------------------------------------------------
    story.append(Paragraph("DEPLOYMENT TROUBLESHOOTING", badge_style))
    story.append(Paragraph("Resolved Deployment Challenges & Free Tier Behavior", title_style))

    d1 = [
        Paragraph("<b>Deployment Challenges Resolved</b>", ParagraphStyle('D1', parent=body_style, textColor=COLOR_PRIMARY_DARK)),
        Paragraph(bold_b("Hardcoded Localhost URLs", "Replaced <code>http://localhost:5000</code> with <code>VITE_API_URL</code> environment variables in React components."), bullet_style),
        Paragraph(bold_b("Production CORS Errors", "Updated Express backend CORS policy to whitelist Vercel production domain."), bullet_style),
        Paragraph(bold_b("MongoDB Atlas Connection", "Whitelisted <code>0.0.0.0/0</code> in MongoDB Atlas Network Access for Render cloud IPs."), bullet_style),
    ]

    d2 = [
        Paragraph("<b>Known Free Tier Limitations</b>", ParagraphStyle('D2', parent=body_style, textColor=COLOR_PRIMARY_DARK)),
        Paragraph(bold_b("Render Idle Spin-Down", "Render free web services spin down after 15 mins of inactivity. First request after idle takes 30-50s to warm up."), bullet_style),
        Paragraph(bold_b("MongoDB Atlas Storage", "512 MB storage cap on M0 free cluster (more than sufficient for application tracking)."), bullet_style),
        Paragraph(bold_b("Vercel Bandwidth", "100 GB monthly bandwidth limit on free hobby tier."), bullet_style),
    ]

    t_debug = Table([[d1, d2]], colWidths=[430, 430])
    t_debug.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), COLOR_BG_LIGHT),
        ('BACKGROUND', (1,0), (1,0), colors.HexColor('#FEF2F2')),
        ('BOX', (0,0), (0,0), 1, COLOR_BORDER),
        ('BOX', (1,0), (1,0), 1, colors.HexColor('#FCA5A5')),
        ('PADDING', (0,0), (-1,-1), 12),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(t_debug)

    doc.build(story, canvasmaker=W9NumberedCanvas)
    print(f"[SUCCESS] Built PDF: {filename}")


def build_w9_pptx(filename):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    C_PRIMARY = RGBColor(79, 70, 229)
    C_PRIMARY_DARK = RGBColor(55, 48, 163)
    C_DARK = RGBColor(30, 41, 59)
    C_MUTED = RGBColor(100, 116, 139)

    def add_w9_header(slide, title_text):
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_PRIMARY
        top_bar.line.color.rgb = C_PRIMARY

        tx = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(10), Inches(0.4))
        tf = tx.text_frame
        p0 = tf.paragraphs[0]
        p0.text = "WEEK 9 DEPLOYMENT DELIVERABLE".upper()
        p0.font.bold = True
        p0.font.size = Pt(11)
        p0.font.color.rgb = C_PRIMARY

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.bold = True
        p1.font.size = Pt(22)
        p1.font.color.rgb = C_PRIMARY_DARK

        foot = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12.133), Inches(0.4))
        fp = foot.text_frame.paragraphs[0]
        fp.text = f"{PROJECT_TITLE} | {WEEK_TITLE} | Intern ID: {INTERN_ID} | {INTERN_NAME}"
        fp.font.size = Pt(10)
        fp.font.color.rgb = C_MUTED

    # Slide 1: Title
    s1 = prs.slides.add_slide(blank_layout)
    sb = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.3), Inches(7.5))
    sb.fill.solid()
    sb.fill.fore_color.rgb = C_PRIMARY
    sb.line.color.rgb = C_PRIMARY

    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "WEEK 9 SUBMISSION DELIVERABLE"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_PRIMARY

    p = tf1.add_paragraph()
    p.text = PROJECT_TITLE
    p.font.bold = True
    p.font.size = Pt(40)
    p.font.color.rgb = C_PRIMARY_DARK

    p = tf1.add_paragraph()
    p.text = WEEK_TITLE
    p.font.size = Pt(18)
    p.font.color.rgb = C_MUTED
    p.space_after = Pt(24)

    p = tf1.add_paragraph()
    p.text = f"Intern Name: {INTERN_NAME}\nIntern ID: {INTERN_ID}\nUniversity: {UNIVERSITY}\nLive Frontend: {URL_FRONTEND}\nLive Backend: {URL_BACKEND}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_DARK

    # Slide 2: Deployment Architecture
    s2 = prs.slides.add_slide(blank_layout)
    add_w9_header(s2, "Production Cloud Hosting Architecture")
    tb2 = s2.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.133), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    pts = [
        ("Frontend Deployment (Vercel)", "Deployed from GitHub root. Configured environment variables (VITE_API_URL) for production backend connection."),
        ("Backend Deployment (Render)", "Node.js Express web service running npm start. Configured environment secrets: MONGO_URI, JWT_SECRET, FRONTEND_ORIGIN."),
        ("Database Hosting (MongoDB Atlas)", "Cloud-hosted M0 Free Cluster with Mongoose ODM and whitelisted IP network access (0.0.0.0/0)."),
        ("Security & Auth", "Stateless JWT authentication tokens and rate-limited endpoints.")
    ]
    for title, desc in pts:
        p = tf2.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(14)
        p.font.color.rgb = C_DARK
        p.space_after = Pt(12)

    # Slide 3: Live Application Screenshots
    s3 = prs.slides.add_slide(blank_layout)
    add_w9_header(s3, "Live Production Verification & Screenshots")
    if os.path.exists(IMG_DASHBOARD):
        s3.shapes.add_picture(IMG_DASHBOARD, Inches(0.6), Inches(1.8), width=Inches(5.8))
    if os.path.exists(IMG_MODAL):
        s3.shapes.add_picture(IMG_MODAL, Inches(6.8), Inches(1.8), width=Inches(5.8))

    # Slide 4: Debugging & Limitations
    s4 = prs.slides.add_slide(blank_layout)
    add_w9_header(s4, "Deployment Debugging & Free Tier Limitations")
    tb4 = s4.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.133), Inches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    pts_debug = [
        ("CORS Resolution", "Updated Express backend CORS policy to dynamically match Vercel production origin."),
        ("Localhost URL Removal", "Replaced hardcoded localhost URLs across AuthContext, Dashboard, and Login components."),
        ("Render Idle Spin-Down", "Render free tier spins down after 15 mins of inactivity. First request takes 30-50 seconds to warm up."),
        ("MongoDB Atlas Whitelisting", "Whitelisted 0.0.0.0/0 in Atlas Network Access for Render cloud IPs.")
    ]
    for title, desc in pts_debug:
        p = tf4.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(14)
        p.font.color.rgb = C_DARK
        p.space_after = Pt(12)

    prs.save(filename)
    print(f"[SUCCESS] Built PPTX: {filename}")


if __name__ == '__main__':
    build_w9_pdf(f"W9_DeploymentProof_{INTERN_ID}.pdf")
    build_w9_pdf(f"W9_FinalPPT_{INTERN_ID}.pdf")
    build_w9_pptx(f"W9_FinalPPT_{INTERN_ID}.pptx")
